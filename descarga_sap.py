"""Descarga flexible de anexos OC/HES. Ejecutar en Windows con SAP GUI abierto.

Dependencias: pip install pywin32 pywinauto openpyxl
La vista previa funciona sin conectarse a SAP.
"""
from __future__ import annotations

import csv
import os
import queue
import re
import threading
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import openpyxl
from openpyxl.utils import column_index_from_string


@dataclass
class Item:
    tipo: str
    codigo: str
    bloque: str = ""
    filas: list[int] = field(default_factory=list)
    oc_padre: str = ""


def codigo(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def codigos(valor) -> list[str]:
    return [v.strip() for v in re.split(r"[,;\n]+", codigo(valor)) if v.strip()]


def carpeta_segura(valor: str) -> str:
    """Impide rutas relativas/absolutas provenientes de celdas Excel."""
    limpio = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", valor).strip(" .")
    if not limpio or limpio.upper() in {"CON", "PRN", "AUX", "NUL"} or re.fullmatch(r"(?:COM|LPT)[1-9]", limpio.upper()):
        limpio = "_" + limpio
    return limpio[:100]


def columna(texto: str, obligatoria: bool) -> int | None:
    texto = texto.strip().upper()
    if not texto and not obligatoria:
        return None
    try:
        return column_index_from_string(texto)
    except (ValueError, TypeError):
        raise ValueError(f"Columna Excel no válida: {texto or '(vacía)'}") from None


def leer_excel(ruta: str, hoja: str, inicio: int, col_oc: str, col_hes: str,
               col_bloque: str, modo: str) -> list[Item]:
    if inicio < 1:
        raise ValueError("La fila inicial debe ser 1 o mayor.")
    if not Path(ruta).is_file():
        raise ValueError("Selecciona un archivo Excel existente.")
    if Path(ruta).suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Se admiten archivos .xlsx y .xlsm. Convierte los .xls a .xlsx.")
    oc_idx = columna(col_oc, modo in {"OC", "OC + HES"})
    hes_idx = columna(col_hes, modo in {"HES", "OC + HES"})
    bloque_idx = columna(col_bloque, False)
    if modo == "HES":
        oc_idx = None
    if modo == "OC":
        hes_idx = None
    wb = openpyxl.load_workbook(ruta, read_only=False, data_only=True)
    try:
        if hoja not in wb:
            raise ValueError(f"Hoja '{hoja}' no encontrada. Disponibles: {', '.join(wb.sheetnames)}")
        ws = wb[hoja]
        # Solo expandimos combinaciones de las columnas utilizadas.
        relevantes = {c for c in (oc_idx, hes_idx, bloque_idx) if c}
        merged = {}
        for rango in ws.merged_cells.ranges:
            for col in relevantes:
                if rango.min_col <= col <= rango.max_col:
                    ancla = ws.cell(rango.min_row, rango.min_col).value
                    for fila in range(rango.min_row, rango.max_row + 1):
                        merged[(fila, col)] = ancla

        def valor(fila, col):
            return merged.get((fila, col), ws.cell(fila, col).value) if col else None

        resultado: dict[tuple[str, str, str], Item] = {}
        bloque_actual = ""
        oc_actual = ""
        for fila in range(inicio, ws.max_row + 1):
            bloque_celda = codigo(valor(fila, bloque_idx))
            if bloque_celda:
                bloque_actual = bloque_celda
            bloque = bloque_actual if bloque_idx else ""
            oc_celda = codigo(valor(fila, oc_idx))
            if oc_celda:
                oc_actual = oc_celda
            oc = oc_actual if modo == "OC + HES" else oc_celda
            hes = codigos(valor(fila, hes_idx))
            if modo in {"OC", "OC + HES"} and oc_celda:
                clave = (bloque, "OC", oc_celda)
                item = resultado.setdefault(clave, Item("OC", oc_celda, bloque))
                if fila not in item.filas:
                    item.filas.append(fila)
            if modo in {"HES", "OC + HES"}:
                for h in hes:
                    # En modo conjunto, la HES queda vinculada a su OC por la clave.
                    clave = (bloque, f"HES:{oc}" if modo == "OC + HES" else "HES", h)
                    item = resultado.setdefault(clave, Item("HES", h, bloque, oc_padre=oc if modo == "OC + HES" else ""))
                    if fila not in item.filas:
                        item.filas.append(fila)
        if not resultado:
            raise ValueError("No se encontraron códigos en las columnas y filas seleccionadas.")
        return list(resultado.values())
    finally:
        wb.close()


def leer_manual(texto: str, modo: str) -> list[Item]:
    """Una línea por documento: OC 450..., HES 100... o código simple."""
    items: dict[tuple[str, str], Item] = {}
    for n, linea in enumerate(texto.splitlines(), 1):
        linea = linea.strip()
        if not linea:
            continue
        match = re.fullmatch(r"(?:(OC|HES)\s*[:;,-]?\s*)?([^\s,;]+)", linea, re.I)
        if not match:
            raise ValueError(f"Línea {n}: usa 'OC 4500001234' o 'HES 10001234'.")
        tipo = (match.group(1) or modo).upper()
        if tipo not in {"OC", "HES"}:
            raise ValueError(f"Línea {n}: indica OC o HES delante del código.")
        clave = (tipo, match.group(2))
        items.setdefault(clave, Item(*clave)).filas.append(n)
    if not items:
        raise ValueError("Escribe al menos un código.")
    return list(items.values())


def destino_item(base: Path, item: Item) -> Path:
    ruta = base / carpeta_segura(item.bloque) if item.bloque else base
    if item.oc_padre:
        return ruta / "OC" / carpeta_segura(item.oc_padre) / "HES" / carpeta_segura(item.codigo)
    return ruta / item.tipo / carpeta_segura(item.codigo)


def tiene_archivos(ruta: Path) -> bool:
    return ruta.is_dir() and any(p.is_file() and p.stat().st_size > 0 for p in ruta.iterdir())


def ejecutar(items: list[Item], base: Path, reanudar: bool, pausa: float,
            detener: threading.Event, avisar):
    import pythoncom
    import sap_backend as sap

    pythoncom.CoInitialize()
    try:
        session = sap.get_sap_session()
        if not sap._sesion_esta_viva(session):
            raise RuntimeError("La sesión SAP no responde. Abre SAP Logon e inicia sesión.")
        registros = []
        for indice, item in enumerate(items, 1):
            if detener.is_set():
                break
            ruta = destino_item(base, item)
            estado, detalle = "", ""
            avisar("log", f"[{indice}/{len(items)}] {item.tipo} {item.codigo}")
            if reanudar and tiene_archivos(ruta):
                estado = "Omitido: carpeta con archivos"
            else:
                try:
                    ruta.mkdir(parents=True, exist_ok=True)
                    if item.tipo == "OC":
                        sap.abrir_oc(session, item.codigo, lambda m: avisar("log", m))
                    else:
                        sap.abrir_hes(session, item.codigo, lambda m: avisar("log", m))
                    if sap.abrir_lista_anexos(session, lambda m: avisar("log", m)):
                        comentario = sap.descargar_todos_los_anexos(session, str(ruta), lambda m: avisar("log", m))
                        if comentario:
                            detalle = "La lista contiene comentario SAP"
                    estado = "Con archivos" if tiene_archivos(ruta) else "Sin archivos"
                except Exception as exc:
                    estado, detalle = "Error", str(exc)
                    avisar("log", f"Error en {item.tipo} {item.codigo}: {exc}")
                finally:
                    try:
                        sap._recuperar_sesion(session, lambda m: avisar("log", m))
                        sap.volver_pantalla_inicial(session, lambda m: avisar("log", m))
                    except Exception as exc:
                        avisar("log", f"SAP requiere revisión: {exc}")
            registros.append([datetime.now().isoformat(timespec="seconds"), item.bloque,
                              item.tipo, item.codigo, ",".join(map(str, item.filas)),
                              estado, detalle, str(ruta)])
            guardar_reporte(base, registros)
            avisar("progreso", (indice, len(items)))
            if pausa and indice < len(items) and not detener.is_set():
                detener.wait(pausa)
        avisar("fin", f"Proceso {'detenido' if detener.is_set() else 'finalizado'}. "
              f"{len(registros)} documentos registrados. Reporte: {base / 'reporte_descargas.csv'}")
    except Exception as exc:
        avisar("fin", f"No se pudo continuar: {exc}")
    finally:
        pythoncom.CoUninitialize()


def guardar_reporte(base: Path, registros: list[list[str]]) -> None:
    base.mkdir(parents=True, exist_ok=True)
    ruta = base / "reporte_descargas.csv"
    temporal = base / "reporte_descargas.csv.tmp"
    with temporal.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["fecha", "bloque", "tipo", "codigo", "filas_origen", "estado", "detalle", "carpeta"])
        writer.writerows(registros)
    os.replace(temporal, ruta)


class App:
    def __init__(self, root: tk.Tk):
        self.root = root
        root.title("Descarga de anexos SAP · OC y HES")
        root.geometry("900x780")
        root.minsize(760, 640)
        self.eventos = queue.Queue()
        self.detener = threading.Event()
        self.items: list[Item] = []
        self.fuente = tk.StringVar(value="Excel")
        self.modo = tk.StringVar(value="OC + HES")
        self.excel = tk.StringVar()
        self.destino = tk.StringVar()
        self.hoja = tk.StringVar(value="Plan de Descarga")
        self.fila = tk.StringVar(value="2")
        self.col_bloque = tk.StringVar(value="")
        self.col_oc = tk.StringVar(value="D")
        self.col_hes = tk.StringVar(value="H")
        self.reanudar = tk.BooleanVar(value=True)
        self.pausa = tk.StringVar(value="0")
        self.estado = tk.StringVar(value="Configura la fuente y revisa la vista previa.")
        self._crear_ui()
        self.root.after(150, self._vaciar_eventos)

    def _crear_ui(self):
        marco = ttk.Frame(self.root, padding=16)
        marco.pack(fill="both", expand=True)
        ttk.Label(marco, text="Descarga de anexos SAP", font=("Segoe UI", 18, "bold")).pack(anchor="w")
        ttk.Label(marco, text="1. Elige datos  ·  2. Revisa la lista  ·  3. Descarga", foreground="#555").pack(anchor="w", pady=(0, 12))
        opciones = ttk.LabelFrame(marco, text="Documentos y origen", padding=10)
        opciones.pack(fill="x")
        linea = ttk.Frame(opciones)
        linea.pack(fill="x")
        ttk.Label(linea, text="Descargar").pack(side="left")
        modo = ttk.Combobox(linea, textvariable=self.modo, values=["OC", "HES", "OC + HES"], state="readonly", width=12)
        modo.pack(side="left", padx=(8, 24))
        ttk.Radiobutton(linea, text="Desde Excel", variable=self.fuente, value="Excel", command=self._alternar).pack(side="left")
        ttk.Radiobutton(linea, text="Escribir códigos", variable=self.fuente, value="Manual", command=self._alternar).pack(side="left", padx=12)
        self.excel_panel = ttk.Frame(opciones)
        self.excel_panel.pack(fill="x", pady=(10, 0))
        self._ruta(self.excel_panel, "Archivo Excel", self.excel, self._buscar_excel)
        columnas = ttk.Frame(self.excel_panel)
        columnas.pack(fill="x", pady=5)
        for titulo, var, ancho in [("Hoja", self.hoja, 18), ("Fila inicial", self.fila, 5), ("Col. OC", self.col_oc, 5), ("Col. HES", self.col_hes, 5), ("Col. bloque opcional", self.col_bloque, 5)]:
            ttk.Label(columnas, text=titulo).pack(side="left", padx=(0, 4))
            ttk.Entry(columnas, textvariable=var, width=ancho).pack(side="left", padx=(0, 14))
        ttk.Label(self.excel_panel, text="Deja Bloque vacío si no lo usas. Se ignoran las columnas que no corresponden al tipo elegido.", foreground="#555").pack(anchor="w")
        self.manual_panel = ttk.Frame(opciones)
        ttk.Label(self.manual_panel, text="Un documento por línea. Para mezclar tipos: OC 4500001234 / HES 10001234").pack(anchor="w")
        self.manual = tk.Text(self.manual_panel, height=5, font=("Consolas", 10))
        self.manual.pack(fill="x")
        salida = ttk.LabelFrame(marco, text="Destino y ejecución", padding=10)
        salida.pack(fill="x", pady=10)
        self._ruta(salida, "Carpeta destino", self.destino, self._buscar_destino)
        controles = ttk.Frame(salida)
        controles.pack(fill="x", pady=(6, 0))
        ttk.Checkbutton(controles, text="Omitir documentos cuya carpeta ya tenga archivos", variable=self.reanudar).pack(side="left")
        ttk.Label(controles, text="Pausa (segundos)").pack(side="left", padx=(24, 5))
        ttk.Entry(controles, textvariable=self.pausa, width=6).pack(side="left")
        botones = ttk.Frame(marco)
        botones.pack(fill="x", pady=5)
        ttk.Button(botones, text="Revisar lista", command=self.previsualizar).pack(side="left")
        self.iniciar = ttk.Button(botones, text="Iniciar descarga", command=self._iniciar, state="disabled")
        self.iniciar.pack(side="left", padx=8)
        self.parar = ttk.Button(botones, text="Detener", command=self._detener, state="disabled")
        self.parar.pack(side="left")
        ttk.Label(marco, textvariable=self.estado).pack(anchor="w", pady=(7, 4))
        self.barra = ttk.Progressbar(marco, mode="determinate")
        self.barra.pack(fill="x")
        tabla = ttk.LabelFrame(marco, text="Vista previa", padding=5)
        tabla.pack(fill="both", expand=True, pady=10)
        self.tabla = ttk.Treeview(tabla, columns=("tipo", "codigo", "bloque", "filas"), show="headings", height=8)
        for col, titulo, ancho in [("tipo", "Tipo", 70), ("codigo", "Código", 180), ("bloque", "Bloque", 140), ("filas", "Filas de origen", 160)]:
            self.tabla.heading(col, text=titulo)
            self.tabla.column(col, width=ancho)
        self.tabla.pack(fill="both", expand=True)
        registro = ttk.LabelFrame(marco, text="Actividad", padding=5)
        registro.pack(fill="both", expand=True)
        self.log = tk.Text(registro, height=8, state="disabled", font=("Consolas", 9))
        self.log.pack(fill="both", expand=True)

    def _ruta(self, padre, titulo, var, accion):
        fila = ttk.Frame(padre)
        fila.pack(fill="x", pady=3)
        ttk.Label(fila, text=titulo, width=15).pack(side="left")
        ttk.Entry(fila, textvariable=var).pack(side="left", fill="x", expand=True, padx=6)
        ttk.Button(fila, text="Buscar…", command=accion).pack(side="left")

    def _alternar(self):
        self.items = []
        self.iniciar.configure(state="disabled")
        for panel in (self.excel_panel, self.manual_panel):
            panel.pack_forget()
        if self.fuente.get() == "Excel":
            self.excel_panel.pack(fill="x", pady=(10, 0))
        else:
            self.manual_panel.pack(fill="x", pady=(10, 0))

    def _buscar_excel(self):
        ruta = filedialog.askopenfilename(filetypes=[("Excel moderno", "*.xlsx *.xlsm")])
        if ruta:
            self.excel.set(ruta)

    def _buscar_destino(self):
        ruta = filedialog.askdirectory()
        if ruta:
            self.destino.set(ruta)

    def previsualizar(self):
        try:
            if self.fuente.get() == "Excel":
                items = leer_excel(self.excel.get().strip(), self.hoja.get().strip(),
                                   int(self.fila.get()), self.col_oc.get(), self.col_hes.get(),
                                   self.col_bloque.get(), self.modo.get())
            else:
                items = leer_manual(self.manual.get("1.0", "end"), self.modo.get())
            destinos = {}
            for item in items:
                ruta = str(destino_item(Path("destino"), item)).casefold()
                if ruta in destinos and destinos[ruta] != (item.tipo, item.codigo, item.bloque):
                    raise ValueError(f"Dos códigos producirían la misma carpeta: {item.codigo}.")
                destinos[ruta] = (item.tipo, item.codigo, item.bloque)
            self.items = items
            self.tabla.delete(*self.tabla.get_children())
            for item in items:
                self.tabla.insert("", "end", values=(item.tipo, item.codigo, item.bloque, ", ".join(map(str, item.filas))))
            self.estado.set(f"{len(items)} documentos listos para descargar.")
            self.iniciar.configure(state="normal")
        except Exception as exc:
            self.items = []
            self.iniciar.configure(state="disabled")
            messagebox.showerror("Revisa los datos", str(exc))

    def _iniciar(self):
        # Releer siempre: la configuración puede cambiar después de la vista previa.
        self.previsualizar()
        if not self.items:
            return
        try:
            base = Path(self.destino.get().strip())
            if not self.destino.get().strip():
                raise ValueError("Selecciona una carpeta destino.")
            pausa = float(self.pausa.get())
            if pausa < 0:
                raise ValueError("La pausa no puede ser negativa.")
            if not self.items:
                raise ValueError("Revisa la lista antes de iniciar.")
            base.mkdir(parents=True, exist_ok=True)
        except Exception as exc:
            messagebox.showerror("Revisa la configuración", str(exc))
            return
        self.detener.clear()
        self.iniciar.configure(state="disabled")
        self.parar.configure(state="normal")
        self.barra.configure(maximum=len(self.items), value=0)
        self.estado.set("Conectando con SAP…")
        threading.Thread(target=ejecutar, args=(list(self.items), base, self.reanudar.get(),
                         pausa, self.detener, self._avisar), daemon=True).start()

    def _detener(self):
        self.detener.set()
        self.parar.configure(state="disabled")
        self.estado.set("Deteniendo al terminar el documento actual…")

    def _avisar(self, tipo, dato):
        self.eventos.put((tipo, dato))

    def _vaciar_eventos(self):
        try:
            while True:
                tipo, dato = self.eventos.get_nowait()
                if tipo == "log":
                    self.log.configure(state="normal")
                    self.log.insert("end", f"[{datetime.now():%H:%M:%S}] {dato}\n")
                    self.log.see("end")
                    self.log.configure(state="disabled")
                elif tipo == "progreso":
                    self.barra.configure(value=dato[0])
                    self.estado.set(f"Procesados {dato[0]} de {dato[1]}")
                elif tipo == "fin":
                    self.estado.set(dato)
                    self.parar.configure(state="disabled")
                    self.iniciar.configure(state="normal")
                    messagebox.showinfo("Descarga SAP", dato)
        except queue.Empty:
            pass
        self.root.after(150, self._vaciar_eventos)


if __name__ == "__main__":
    ventana = tk.Tk()
    App(ventana)
    ventana.mainloop()
